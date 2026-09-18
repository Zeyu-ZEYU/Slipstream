"""Tiny dual renderer for the paper's block diagrams and timelines.

One shape list renders twice: to a camera-ready PDF/PNG through matplotlib and
to an editable PPTX through python-pptx (the same geometry, in inches, origin at
the top-left corner). Edit the PPTX in PowerPoint and export to PDF when a hand
adjustment is needed; this machine has no PPTX-to-PDF converter.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Nimbus Sans"      # Helvetica clone available to matplotlib here
PPTX_FONT = "Arial"       # its PowerPoint counterpart
DASH = (0, (3, 2))


def _rgb(hexcolor):
    return RGBColor.from_string(hexcolor.lstrip("#"))


class Canvas:
    def __init__(self, w, h):
        self.w, self.h, self.items = w, h, []

    # ---- shapes -------------------------------------------------------------
    def box(self, x, y, w, h, text="", fill="#FFFFFF", line="#000000", lw=0.8, fs=8,
            color="#000000", bold=False, italic=False, radius=0.05, dash=False,
            align="center", valign="center", pad=0.03):
        self.items.append(("box", dict(x=x, y=y, w=w, h=h, text=text, fill=fill, line=line,
                                       lw=lw, fs=fs, color=color, bold=bold, italic=italic,
                                       radius=radius, dash=dash, align=align, valign=valign,
                                       pad=pad)))

    def text(self, x, y, w, h, text, fs=8, color="#000000", align="center", valign="center",
             bold=False, italic=False, rotation=0):
        self.items.append(("text", dict(x=x, y=y, w=w, h=h, text=text, fs=fs, color=color,
                                        align=align, valign=valign, bold=bold, italic=italic,
                                        rotation=rotation)))

    def arrow(self, x1, y1, x2, y2, color="#333333", lw=1.0, head=True, dash=False, tail=False):
        self.items.append(("arrow", dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, lw=lw,
                                         head=head, dash=dash, tail=tail)))

    def line(self, x1, y1, x2, y2, color="#333333", lw=0.8, dash=False):
        self.arrow(x1, y1, x2, y2, color=color, lw=lw, head=False, dash=dash)

    def group_begin(self):
        """Start a PowerPoint group; matplotlib ignores it."""
        self.items.append(("group_begin", {}))

    def group_end(self):
        self.items.append(("group_end", {}))

    def badge(self, cx, cy, text, r=0.07, fs=8, fill="#FFFFFF", line="#333333", color="#333333", lw=0.7):
        """Numbered circle centered at (cx, cy)."""
        self.items.append(("badge", dict(cx=cx, cy=cy, r=r, text=text, fs=fs, fill=fill, line=line,
                                         color=color, lw=lw)))

    # ---- matplotlib ---------------------------------------------------------
    def _mpl(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
        plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": [FONT, "DejaVu Sans"],
                             "pdf.fonttype": 42, "ps.fonttype": 42})
        fig = plt.figure(figsize=(self.w, self.h))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, self.w); ax.set_ylim(self.h, 0); ax.axis("off")
        ha = {"center": "center", "left": "left", "right": "right"}
        for kind, d in self.items:
            if kind in ("group_begin", "group_end"):
                continue
            if kind == "box":
                r = min(d["radius"], d["w"] / 2, d["h"] / 2)
                patch = FancyBboxPatch((d["x"] + r, d["y"] + r), d["w"] - 2 * r, d["h"] - 2 * r,
                                       boxstyle=f"round,pad={r},rounding_size={r}",
                                       facecolor=d["fill"], edgecolor=d["line"],
                                       linewidth=d["lw"], linestyle=DASH if d["dash"] else "-",
                                       mutation_aspect=1)
                ax.add_patch(patch)
                if d["text"]:
                    self._mpl_text(ax, d, ha)
            elif kind == "text":
                self._mpl_text(ax, d, ha)
            elif kind == "badge":
                from matplotlib.patches import Circle
                ax.add_patch(Circle((d["cx"], d["cy"]), d["r"], facecolor=d["fill"], edgecolor=d["line"],
                                    linewidth=d["lw"], zorder=5))
                ax.text(d["cx"], d["cy"], d["text"], fontsize=d["fs"], color=d["color"], ha="center",
                        va="center", zorder=6, fontweight="bold")
            elif kind == "arrow":
                style = "-|>" if d["head"] else "-"
                if d["tail"]:
                    style = "<|-|>"
                p = FancyArrowPatch((d["x1"], d["y1"]), (d["x2"], d["y2"]), arrowstyle=style,
                                    mutation_scale=7 + 3 * d["lw"], linewidth=d["lw"],
                                    color=d["color"], linestyle=DASH if d["dash"] else "-",
                                    shrinkA=0, shrinkB=0)
                ax.add_patch(p)
        return fig

    @staticmethod
    def _mpl_text(ax, d, ha):
        al = d["align"]; va = d["valign"]
        tx = d["x"] + (d["w"] / 2 if al == "center" else (d.get("pad", 0.03) if al == "left" else d["w"] - d.get("pad", 0.03)))
        ty = d["y"] + (d["h"] / 2 if va == "center" else (d.get("pad", 0.03) if va == "top" else d["h"] - d.get("pad", 0.03)))
        ax.text(tx, ty, d["text"], fontsize=d["fs"], color=d["color"], ha=ha[al],
                va={"center": "center", "top": "top", "bottom": "bottom"}[va],
                fontweight="bold" if d["bold"] else "normal",
                fontstyle="italic" if d["italic"] else "normal", linespacing=1.15,
                rotation=d.get("rotation", 0), multialignment=al)

    # ---- python-pptx --------------------------------------------------------
    def _pptx(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(max(self.w, 1.0)), Inches(max(self.h, 1.0))   # PowerPoint needs at least 1 in
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        palign = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}
        vanchor = {"center": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}
        container = slide.shapes
        for kind, d in self.items:
            if kind == "group_begin":
                container = slide.shapes.add_group_shape().shapes
                continue
            if kind == "group_end":
                container = slide.shapes
                continue
            if kind == "box":
                sp = container.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(d["x"]), Inches(d["y"]),
                                            Inches(d["w"]), Inches(d["h"]))
                sp.adjustments[0] = min(0.5, d["radius"] / max(min(d["w"], d["h"]), 1e-6))
                sp.fill.solid(); sp.fill.fore_color.rgb = _rgb(d["fill"])
                sp.line.color.rgb = _rgb(d["line"]); sp.line.width = Pt(d["lw"])
                sp.shadow.inherit = False
                if d["dash"]:
                    ln = sp.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
                self._pptx_text(sp.text_frame, d, palign, vanchor)
            elif kind == "text":
                tb = container.add_textbox(Inches(d["x"]), Inches(d["y"]), Inches(d["w"]), Inches(d["h"]))
                tb.rotation = -d.get("rotation", 0)
                self._pptx_text(tb.text_frame, d, palign, vanchor, margins=True)
            elif kind == "badge":
                sp = container.add_shape(MSO_SHAPE.OVAL, Inches(d["cx"] - d["r"]), Inches(d["cy"] - d["r"]),
                                            Inches(2 * d["r"]), Inches(2 * d["r"]))
                sp.fill.solid(); sp.fill.fore_color.rgb = _rgb(d["fill"])
                sp.line.color.rgb = _rgb(d["line"]); sp.line.width = Pt(d["lw"]); sp.shadow.inherit = False
                tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                pgh = tf.paragraphs[0]; pgh.alignment = PP_ALIGN.CENTER
                run = pgh.add_run(); run.text = d["text"]; run.font.size = Pt(d["fs"]); run.font.bold = True
                run.font.name = PPTX_FONT; run.font.color.rgb = _rgb(d["color"])
            elif kind == "arrow":
                c = container.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(d["x1"]), Inches(d["y1"]),
                                               Inches(d["x2"]), Inches(d["y2"]))
                c.line.color.rgb = _rgb(d["color"]); c.line.width = Pt(d["lw"])
                ln = c.line._get_or_add_ln()
                if d["head"]:
                    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
                if d["tail"]:
                    ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
                if d["dash"]:
                    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
        return prs

    @staticmethod
    def _pptx_text(tf, d, palign, vanchor, margins=False):
        tf.word_wrap = True; tf.vertical_anchor = vanchor[d["valign"]]
        m = Inches(d.get("pad", 0.03)) if margins or True else 0
        tf.margin_left = tf.margin_right = m; tf.margin_top = tf.margin_bottom = Inches(0.01)
        lines = d["text"].split("\n") if d["text"] else [""]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = palign[d["align"]]
            r = p.add_run(); r.text = ln
            r.font.size = Pt(d["fs"]); r.font.name = PPTX_FONT; r.font.bold = d["bold"]
            r.font.italic = d["italic"]; r.font.color.rgb = _rgb(d["color"])

    # ---- output -------------------------------------------------------------
    def save(self, stem, dpi=300):
        fig = self._mpl()
        fig.savefig(f"{stem}.pdf"); fig.savefig(f"{stem}.png", dpi=dpi)
        self._pptx().save(f"{stem}.pptx")
        print("wrote", f"{stem}.pdf", f"{stem}.png", f"{stem}.pptx")
